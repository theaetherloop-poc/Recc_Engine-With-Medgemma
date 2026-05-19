"""
Examine.com Condition Coverage Evaluation

For each condition in NatMed_Pro_Conditions.xlsx, queries the Apify
Examine.com scraper and records how many supplements Examine recommends.

Run:    .venv/bin/python evaluate_examine_coverage.py
Resume: re-run — already-processed conditions are skipped.

Output:
  test_outputs/examine_coverage/progress.json          ← live progress (resumable)
  test_outputs/examine_coverage/examine_coverage.xlsx  ← final report
"""

import os
import json
import time
from pathlib import Path
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

APIFY_API_TOKEN = os.getenv("APIFY_API_TOKEN")
ACTOR_ID = "hanamira/examine-com-supplement-research"
OUTPUT_DIR = Path("test_outputs/examine_coverage")
PROGRESS_FILE = OUTPUT_DIR / "progress.json"
EXCEL_OUT = OUTPUT_DIR / "examine_coverage.xlsx"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

if not APIFY_API_TOKEN:
    raise EnvironmentError("APIFY_API_TOKEN not found in .env")

try:
    from apify_client import ApifyClient
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    raise ImportError(f"Missing package: {e}")

client = ApifyClient(APIFY_API_TOKEN)


# ── Load conditions ───────────────────────────────────────────────────────────

def load_conditions(path: str) -> list[str]:
    """Return list of condition names from the NatMed Excel, skipping section headers."""
    wb = openpyxl.load_workbook(path)
    ws = wb["NatMed Pro Conditions"]
    conditions = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        name = str(row[0]).strip() if row[0] else None
        if not name:
            continue
        vals = row[1:]
        # Section headers: long parenthetical names with all-None numeric columns
        is_header = all(v is None for v in vals) and "(" in name and len(name) > 40
        if not is_header:
            conditions.append(name)
    return conditions


# ── Match quality ─────────────────────────────────────────────────────────────

def _match_quality(search_term: str, found_name: str) -> str:
    """Classify match: 'exact', 'partial', or 'unrelated'."""
    if not found_name:
        return "unrelated"
    s, f = search_term.lower(), found_name.lower()
    if s == f or s in f or f in s:
        return "exact"
    s_words = {w for w in s.split() if len(w) > 4}
    f_words = {w for w in f.split() if len(w) > 4}
    if s_words & f_words:
        return "partial"
    return "unrelated"


# ── Apify query ───────────────────────────────────────────────────────────────

def query_examine(condition_name: str) -> dict:
    """Search Examine.com for a condition and return structured result."""
    try:
        run = client.actor(ACTOR_ID).call(
            run_input={"query": condition_name, "searchType": "conditions", "maxResults": 5}
        )
        dataset_id = run["defaultDatasetId"]
        items = list(client.dataset(dataset_id).iterate_items())

        SKIP = {"Supplements", "Health Conditions and Goals"}
        records = [
            r for item in items
            for r in item.get("results", [])
            if r.get("type") == "condition" and r.get("name") not in SKIP
        ]
        all_names = [r.get("name") for r in records]

        if not records:
            return {
                "found": False, "condition_name": None, "match_quality": None,
                "supplement_count": 0, "supplements": [], "url": None,
                "all_conditions_found": [], "dataset_id": dataset_id, "error": None,
            }

        # Pick best match (prefer exact > partial > unrelated)
        best, best_q = records[0], "unrelated"
        for r in records:
            q = _match_quality(condition_name, r.get("name", ""))
            if q == "exact":
                best, best_q = r, "exact"
                break
            elif q == "partial" and best_q == "unrelated":
                best, best_q = r, "partial"

        supps = [s["name"] for s in best.get("recommendedSupplements", []) if s.get("name") not in SKIP]

        return {
            "found": True,
            "condition_name": best.get("name"),
            "match_quality": best_q,
            "supplement_count": len(supps),
            "supplements": supps,
            "url": best.get("url"),
            "all_conditions_found": all_names,
            "dataset_id": dataset_id,
            "error": None,
        }

    except Exception as e:
        return {
            "found": False, "condition_name": None, "match_quality": None,
            "supplement_count": 0, "supplements": [], "url": None,
            "all_conditions_found": [], "dataset_id": None, "error": str(e),
        }


# ── Progress helpers ──────────────────────────────────────────────────────────

def load_progress() -> dict:
    if PROGRESS_FILE.exists():
        with open(PROGRESS_FILE) as f:
            return json.load(f)
    return {}


def save_progress(progress: dict):
    with open(PROGRESS_FILE, "w") as f:
        json.dump(progress, f, indent=2)


# ── Excel report ──────────────────────────────────────────────────────────────

def write_excel(conditions: list[str], progress: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Examine Coverage"

    headers = [
        "Condition (searched)",
        "Found on Examine?",
        "Match Quality",
        "Examine Condition Name",
        "Supplement Count",
        "Supplement List",
        "All Conditions Found by Search",
        "Examine URL",
        "Error",
    ]

    hdr_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 28

    green  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red    = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    grey   = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for row_i, cond in enumerate(conditions, 2):
        r = progress.get(cond, {})
        found = r.get("found", False)
        mq = r.get("match_quality")

        row_data = [
            cond,
            "Yes" if found else ("Error" if r.get("error") else "No"),
            mq or "N/A",
            r.get("condition_name") or "",
            r.get("supplement_count") if found else None,
            ", ".join(r.get("supplements", [])),
            ", ".join(r.get("all_conditions_found", [])),
            r.get("url") or "",
            r.get("error") or "",
        ]
        for col, val in enumerate(row_data, 1):
            ws.cell(row=row_i, column=col, value=val)

        # Color: Found? (col 2)
        ws.cell(row=row_i, column=2).fill = green if found else (red if r.get("error") else grey)
        # Color: Match quality (col 3)
        mq_fill = {"exact": green, "partial": yellow, "unrelated": red}.get(mq, grey)
        ws.cell(row=row_i, column=3).fill = mq_fill
        # Color: Supplement count (col 5)
        count = r.get("supplement_count", 0)
        ws.cell(row=row_i, column=5).fill = green if count >= 5 else (yellow if count >= 2 else (red if count == 0 and found else grey))

    col_widths = [55, 14, 14, 35, 14, 70, 50, 45, 30]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    done = {k: v for k, v in progress.items() if k in conditions}
    found_exact   = sum(1 for v in done.values() if v.get("match_quality") == "exact")
    found_partial = sum(1 for v in done.values() if v.get("match_quality") == "partial")
    found_unrel   = sum(1 for v in done.values() if v.get("match_quality") == "unrelated")
    not_found     = sum(1 for v in done.values() if not v.get("found") and not v.get("error"))
    errors        = sum(1 for v in done.values() if v.get("error"))
    counts        = [v.get("supplement_count", 0) for v in done.values() if v.get("found")]

    rows = [
        ("Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total conditions queried", len(conditions)),
        ("Processed so far", len(done)),
        ("", ""),
        ("--- Match Quality ---", ""),
        ("Exact match", found_exact),
        ("Partial match", found_partial),
        ("Unrelated match (wrong condition returned)", found_unrel),
        ("Not found on Examine", not_found),
        ("Errors", errors),
        ("", ""),
        ("--- Supplement Counts (found conditions only) ---", ""),
        ("Avg supplements per found condition", round(sum(counts) / len(counts), 1) if counts else 0),
        ("Max supplements (single condition)", max(counts) if counts else 0),
        ("Min supplements (single condition)", min(counts) if counts else 0),
        ("Conditions with 0 supplements", sum(1 for c in counts if c == 0)),
        ("Conditions with 5+ supplements", sum(1 for c in counts if c >= 5)),
    ]
    for r_i, (label, val) in enumerate(rows, 1):
        ws2.cell(row=r_i, column=1, value=label).font = Font(bold=True if label.startswith("---") or not label else False)
        ws2.cell(row=r_i, column=2, value=val)
    ws2.column_dimensions["A"].width = 48
    ws2.column_dimensions["B"].width = 20

    wb.save(EXCEL_OUT)
    print(f"Excel saved → {EXCEL_OUT}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    conditions = load_conditions("NatMed_Pro_Conditions.xlsx")
    progress = load_progress()

    done = [c for c in conditions if c in progress]
    remaining = [c for c in conditions if c not in progress]

    print(f"Total conditions: {len(conditions)}")
    print(f"Already processed: {len(done)}")
    print(f"Remaining: {len(remaining)}")
    est_min = len(remaining) * 25 // 60
    est_sec = len(remaining) * 25 % 60
    print(f"Estimated time: ~{est_min}m {est_sec}s\n")

    if not remaining:
        print("All done. Regenerating Excel...")
        write_excel(conditions, progress)
        return

    for i, cond in enumerate(remaining):
        print(f"[{len(done) + i + 1}/{len(conditions)}] {cond[:70]}")
        result = query_examine(cond)

        status = "FOUND" if result["found"] else ("ERROR" if result["error"] else "NOT FOUND")
        mq = result.get("match_quality", "")
        exam_name = result.get("condition_name", "")
        count = result.get("supplement_count", 0)
        print(f"  → {status} | match={mq} | examine='{exam_name}' | supplements={count}")
        if result["error"]:
            print(f"  ⚠ {result['error']}")

        progress[cond] = result
        save_progress(progress)

        if i < len(remaining) - 1:
            time.sleep(2)

    print("\nAll done. Writing Excel report...")
    write_excel(conditions, progress)

    # Print final stats
    found_total = sum(1 for v in progress.values() if v.get("found"))
    exact = sum(1 for v in progress.values() if v.get("match_quality") == "exact")
    partial = sum(1 for v in progress.values() if v.get("match_quality") == "partial")
    unrel = sum(1 for v in progress.values() if v.get("match_quality") == "unrelated")
    not_found = sum(1 for v in progress.values() if not v.get("found") and not v.get("error"))
    errors = sum(1 for v in progress.values() if v.get("error"))
    print(f"\nFinal Results ({len(conditions)} conditions):")
    print(f"  Found:      {found_total}  (exact={exact}, partial={partial}, unrelated={unrel})")
    print(f"  Not found:  {not_found}")
    print(f"  Errors:     {errors}")


if __name__ == "__main__":
    main()
